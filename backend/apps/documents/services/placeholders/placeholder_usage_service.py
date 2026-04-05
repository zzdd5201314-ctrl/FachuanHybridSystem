"""Business logic services."""

from __future__ import annotations

import logging
import re
from typing import Any

from django.core.cache import cache
from django.db.models import Count, Max

from apps.core.infrastructure import CacheTimeout
from apps.documents.models import DocumentTemplate

logger = logging.getLogger(__name__)

try:
    from docxtpl import DocxTemplate as _DocxTemplate
except Exception:
    logger.exception("操作失败")

    _DocxTemplate = None

try:
    from docx import Document as _DocxDocument
except Exception:
    logger.exception("操作失败")

    _DocxDocument = None  # type: ignore[assignment]


class PlaceholderUsageService:
    _PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*(.*?)\s*\}\}")
    _CACHE_KEY_PREFIX = "documents:placeholder_usage:"

    def get_usage_map(self) -> dict[str, set[str]]:
        cache_key = self._get_cache_key()
        cached = cache.get(cache_key)
        if cached is not None:
            return {key: set(types) for key, types in cached.items()}

        usage: dict[str, set[str]] = {}
        templates = DocumentTemplate.objects.filter(is_active=True).only(
            "id", "template_type", "file", "file_path", "updated_at", "name"
        )
        for template in templates:
            template_type = (getattr(template, "template_type", "") or "").strip()
            if not template_type:
                continue
            placeholders = self._extract_template_placeholders(template)
            for key in placeholders:
                usage.setdefault(key, set()).add(template_type)

        cache.set(
            cache_key,
            {key: sorted(list(types)) for key, types in usage.items()},
            timeout=CacheTimeout.get_day(),
        )
        return usage

    def _get_cache_key(self) -> str:
        signature = self._get_template_signature()
        return f"{self._CACHE_KEY_PREFIX}{signature[0]}:{signature[1]}:{signature[2]}"

    def _get_template_signature(self) -> tuple[int, int, str]:
        agg = DocumentTemplate.objects.filter(is_active=True).aggregate(
            max_updated_at=Max("updated_at"),
            count=Count("id"),
        )
        max_updated_at = agg.get("max_updated_at")
        count = int(agg.get("count") or 0)
        ts = getattr(max_updated_at, "timestamp", None)
        max_updated_at_ts = int(ts()) if ts else 0
        return max_updated_at_ts, count, "v1"

    def _extract_template_placeholders(self, template: DocumentTemplate) -> set[str]:
        file_path = ""
        try:
            file_path = template.get_file_location()
        except Exception:
            logger.exception(
                "get_template_file_location_failed",
                extra={"template_id": template.pk, "template_name": template.name},
            )
            return set()

        if not file_path:
            return set()

        keys = self._extract_by_docxtpl(file_path)
        if keys is not None:
            return keys

        return self._extract_by_python_docx(file_path)

    def _extract_by_docxtpl(self, file_path: str) -> set[str] | None:
        if _DocxTemplate is None:
            return None

        try:
            doc = _DocxTemplate(file_path)
            keys = set(doc.get_undeclared_template_variables() or [])
            return {str(k).strip() for k in keys if str(k).strip()}
        except Exception:
            logger.exception("extract_placeholders_by_docxtpl_failed", extra={"file_path": file_path})
            return set()

    def _extract_by_python_docx(self, file_path: str) -> set[str]:
        if _DocxDocument is None:
            return set()

        try:
            word_doc = _DocxDocument(file_path)
        except Exception:
            logger.exception("open_docx_failed", extra={"file_path": file_path})
            return set()

        placeholders: set[str] = set()
        for text in self._iter_doc_texts(word_doc):
            for match in self._PLACEHOLDER_PATTERN.findall(text or ""):
                key = (match or "").strip()
                if key:
                    placeholders.add(key)
        return placeholders

    def _iter_doc_texts(self, word_doc: Any) -> list[str]:
        texts: list[str] = []
        try:
            for para in getattr(word_doc, "paragraphs", []) or []:
                texts.append(getattr(para, "text", "") or "")
        except Exception:
            logger.debug("iter_docx_paragraphs_failed", exc_info=True)

        try:
            for table in getattr(word_doc, "tables", []) or []:
                for row in getattr(table, "rows", []) or []:
                    for cell in getattr(row, "cells", []) or []:
                        texts.append(getattr(cell, "text", "") or "")
        except Exception:
            logger.debug("iter_docx_tables_failed", exc_info=True)

        return texts
